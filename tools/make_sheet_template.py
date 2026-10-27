"""구글 시트 양식(.xlsx) 생성. 구글 드라이브에 올린 뒤 'Google 스프레드시트로 열기'로 쓴다.

python tools/make_sheet_template.py  →  docs/구글시트_양식.xlsx
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from contesthub import config, store  # noqa: E402
from contesthub.models import CATEGORIES, SUBCATEGORIES, TARGET_LABELS  # noqa: E402

HEAD_FILL = PatternFill("solid", fgColor="E8EDFF")
HEAD_FONT = Font(bold=True)


def header(ws, columns: list[tuple[str, int, str]]) -> None:
    for i, (name, width, note) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        if note:
            cell.comment = Comment(note, "ContestHub")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"


def dropdown(ws, col: str, options: list[str], prompt: str) -> None:
    # 여러 개를 쉼표로 적을 수 있도록 목록 밖 값도 허용(경고만)
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                        showErrorMessage=False, promptTitle="입력 안내", prompt=prompt, showInputMessage=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}1000")


def main() -> None:
    wb = Workbook()
    targets = list(TARGET_LABELS.values())

    # 사용법
    ws = wb.active
    ws.title = "사용법"
    guide = [
        "ContestHub 운영 시트",
        "",
        "■ 수정 탭: 자동 수집된 공모전을 숨기거나 구분·장르·대상·마감일을 고칩니다.",
        "  - 공모전ID는 사이트 상세 페이지 아래 'ID 복사' 버튼으로 복사합니다.",
        "  - 고칠 칸에만 값을 적고, 나머지는 비워 두면 자동 수집 값을 그대로 씁니다.",
        "  - 여러 개는 쉼표로 적습니다. 예) 웹툰, 게임",
        "■ 직접등록 탭: 자동으로 수집되지 않는 공모전을 추가합니다. 제목과 마감일은 꼭 적어 주세요.",
        "■ 키워드 탭: 자동 분류에 쓰는 키워드입니다. 키워드는 쉼표로 구분합니다.",
        "  - 종류는 구분 / 장르 / 제외 / 대학 / 대학과목 중 하나. '제외'는 제목에 들어 있으면 수집하지 않는 단어입니다.",
        "  - '대학'은 웹툰·게임·애니 관련 학과가 있는 대학 이름(줄임말), '대학과목'은 그 대학 공모전으로 인정할 과목 단어입니다.",
        "    제목에 대학 이름이 있고 과목 단어가 함께 나오면 '대학' 탭에 들어갑니다. ('○○대학교'라고 적힌 제목은 목록에 없어도 인식)",
        "■ 수집사이트 탭: 사이트별로 수집을 켜고(O) 끕니다(X).",
        "",
        "날짜는 2026-11-30 형식으로 적습니다.",
        "수정 내용은 다음 자동 실행(08시·18시)에 반영됩니다. 바로 반영하려면 GitHub Actions에서 'Run workflow'를 누르세요.",
        "반영하지 못한 줄은 사이트의 /review/ 페이지 위쪽에 표시됩니다.",
        "",
        "이 시트는 '웹에 게시(CSV)'로 공개되므로, 공개되어도 괜찮은 내용만 적어 주세요.",
    ]
    for i, line in enumerate(guide, start=1):
        ws.cell(row=i, column=1, value=line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 110

    # 수정
    ws = wb.create_sheet("수정")
    header(ws, [
        ("공모전ID", 14, "사이트 상세 페이지의 'ID 복사'로 복사한 값 (예: 2026-a1b2c3)"),
        ("제목(참고용)", 50, "알아보기 쉽게 적는 칸. 사이트에는 반영되지 않음"),
        ("숨김", 7, "O를 적으면 사이트에서 숨김"),
        ("구분", 14, "웹툰 / 웹소설 / 게임. 여러 개는 쉼표"),
        ("장르", 18, ", ".join(SUBCATEGORIES)),
        ("대상", 22, ", ".join(targets)),
        ("마감일", 12, "예) 2026-11-30"),
        ("제목수정", 40, "사이트에 보일 제목을 바꿀 때만"),
        ("메모", 30, "운영자 메모. 사이트에는 반영되지 않음"),
    ])
    dropdown(ws, "C", ["O"], "숨기려면 O")
    dropdown(ws, "D", list(CATEGORIES), "여러 개는 쉼표로 (예: 웹툰, 게임)")
    dropdown(ws, "E", list(SUBCATEGORIES), "여러 개는 쉼표로")
    dropdown(ws, "F", targets, "여러 개는 쉼표로")
    today = date.today()
    review = [c for c in store.load().values() if c.needs_review and c.status(today) != "마감"]
    for c in sorted(review, key=lambda c: c.end_date or "9999"):
        ws.append([c.id, c.title, "", "", "", "", "", "", "자동: " + ", ".join(c.review_reasons)])

    # 직접등록
    ws = wb.create_sheet("직접등록")
    header(ws, [
        ("제목", 45, "필수"), ("구분", 12, "웹툰 / 웹소설 / 게임"), ("장르", 16, ", ".join(SUBCATEGORIES)),
        ("시작일", 12, "예) 2026-11-01"), ("마감일", 12, "필수. 예) 2026-11-30"), ("주최", 20, ""),
        ("대상", 22, ", ".join(targets)), ("상금", 18, "예) 1,000만원"), ("공식링크", 40, "https://..."),
        ("해외", 7, "해외 공모전이면 O"), ("메모", 30, ""),
    ])
    dropdown(ws, "B", list(CATEGORIES), "여러 개는 쉼표로")
    dropdown(ws, "C", list(SUBCATEGORIES), "여러 개는 쉼표로")
    dropdown(ws, "G", targets, "여러 개는 쉼표로")
    dropdown(ws, "J", ["O"], "해외면 O")

    # 키워드
    ws = wb.create_sheet("키워드")
    header(ws, [("종류", 10, "구분 / 장르 / 제외 / 대학 / 대학과목"), ("이름", 12, "구분·장르 이름. 제외·대학·대학과목은 비워 둠"),
                ("키워드", 110, "쉼표로 구분")])
    dropdown(ws, "A", ["구분", "장르", "제외", "대학", "대학과목"], "구분 / 장르 / 제외 / 대학 / 대학과목")
    kw = config.keywords()
    for name, words in kw["categories"].items():
        ws.append(["구분", name, ", ".join(words)])
    for name, words in kw["subcategories"].items():
        ws.append(["장르", name, ", ".join(words)])
    ws.append(["제외", "", ", ".join(kw["exclude_title"])])
    ws.append(["대학", "", ", ".join(kw.get("universities", []))])
    ws.append(["대학과목", "", ", ".join(kw.get("univ_subjects", []))])
    for row in ws.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True)

    # 수집사이트
    ws = wb.create_sheet("수집사이트")
    header(ws, [("사이트", 14, "바꾸지 마세요"), ("사용", 7, "O: 수집 / X: 중지"), ("설명", 70, "")])
    dropdown(ws, "B", ["O", "X"], "O 또는 X")
    descriptions = {"wevity": "위비티 — 웹툰·게임·문학 분야", "contestkorea": "콘테스트코리아 — 문학·미술/웹툰·IT·아이디어 분야",
                    "itchio": "itch.io — 해외 게임잼 (참가자 200명 이상)",
                    "mgood": "엠굿(월간 미대입시) — 대학 주최 고등학생 웹툰·게임·애니 실기대회·공모전", "thinkcontest": "씽굿 — 보류 (자동 수집 불가)"}
    for site, opts in config.settings()["sites"].items():
        ws.append([site, "O" if opts.get("enabled") else "X", descriptions.get(site, "")])

    out = ROOT / "docs" / "구글시트_양식.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(f"{out} (수정 탭에 확인 필요 {len(review)}건 미리 채움)")


if __name__ == "__main__":
    main()
